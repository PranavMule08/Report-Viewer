import os
import asyncio
import logging
from typing import Tuple, Optional
from pdfplumber import open as pdf_open
from docx import Document
from app.core.config import settings

logger = logging.getLogger(__name__)


class FileExtractionError(Exception):
    """Custom exception for file extraction errors."""
    pass


class FileExtractor:
    """Service for extracting text from various file formats."""

    ALLOWED_EXTENSIONS = settings.ALLOWED_EXTENSIONS
    MAX_FILE_SIZE = settings.MAX_FILE_SIZE_BYTES

    @classmethod
    def validate_file(cls, filename: str, file_size: int) -> Tuple[bool, Optional[str]]:
        """
        Validate uploaded file.
        Returns (is_valid, error_message).
        """
        # Check file size
        if file_size > cls.MAX_FILE_SIZE:
            return False, f"File size exceeds maximum limit of {settings.MAX_FILE_SIZE_MB}MB"

        if file_size == 0:
            return False, "Empty file is not allowed"

        # Check extension
        ext = filename.lower().split('.')[-1] if '.' in filename else ''
        if ext not in cls.ALLOWED_EXTENSIONS:
            return False, f"File type '.{ext}' is not supported. Allowed types: {', '.join(cls.ALLOWED_EXTENSIONS)}"

        return True, None

    @classmethod
    async def extract_text(cls, file_path: str, file_type: str) -> str:
        """
        Extract text from a file based on its type (async interface).
        Returns extracted text or raises FileExtractionError.
        """
        return await asyncio.to_thread(cls.extract_text_sync, file_path, file_type)

    @classmethod
    def extract_text_sync(cls, file_path: str, file_type: str) -> str:
        """
        Synchronous text extraction (safe to run in a thread).
        Returns extracted text or raises FileExtractionError.
        """
        if not os.path.exists(file_path):
            raise FileExtractionError(f"File not found: {file_path}")

        if file_type == "pdf":
            return cls._extract_pdf_sync(file_path)
        elif file_type == "docx":
            return cls._extract_docx_sync(file_path)
        elif file_type == "txt":
            return cls._extract_txt_sync(file_path)
        else:
            raise FileExtractionError(f"Unsupported file type: {file_type}")

    @classmethod
    def _extract_pdf_sync(cls, file_path: str) -> str:
        """Extract text from PDF file."""
        try:
            text_parts = []
            with pdf_open(file_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(f"[Page {i + 1}]\n{page_text}\n")

            full_text = "\n".join(text_parts)

            if not full_text.strip():
                raise FileExtractionError("No text could be extracted from the PDF")

            logger.info(f"Extracted {len(full_text)} characters from PDF")
            return full_text

        except FileExtractionError:
            raise
        except Exception as e:
            logger.error(f"PDF extraction error: {str(e)}")
            raise FileExtractionError(f"Failed to extract text from PDF: {str(e)}")

    @classmethod
    def _extract_docx_sync(cls, file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            doc = Document(file_path)
            paragraphs = []

            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            # Also extract tables if any
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)

            full_text = "\n\n".join(paragraphs)

            if not full_text.strip():
                raise FileExtractionError("No text could be extracted from the DOCX file")

            logger.info(f"Extracted {len(full_text)} characters from DOCX")
            return full_text

        except FileExtractionError:
            raise
        except Exception as e:
            logger.error(f"DOCX extraction error: {str(e)}")
            raise FileExtractionError(f"Failed to extract text from DOCX: {str(e)}")

    @classmethod
    def _extract_txt_sync(cls, file_path: str) -> str:
        """Extract text from TXT file."""
        try:
            # Read raw bytes (sync is fine here - runs in a worker thread)
            with open(file_path, 'rb') as f:
                raw_data = f.read()

            # Check for empty file
            if not raw_data.strip():
                raise FileExtractionError("The text file is empty")

            # Try UTF-8 first, then fallback encodings
            encodings = ['utf-8', 'latin-1', 'cp1252', 'utf-16']
            text = None

            for encoding in encodings:
                try:
                    text = raw_data.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue

            if text is None:
                # Last resort: replace errors
                text = raw_data.decode('utf-8', errors='replace')

            if not text.strip():
                raise FileExtractionError("No readable text found in the file")

            logger.info(f"Extracted {len(text)} characters from TXT")
            return text

        except FileExtractionError:
            raise
        except Exception as e:
            logger.error(f"TXT extraction error: {str(e)}")
            raise FileExtractionError(f"Failed to extract text from TXT: {str(e)}")

    @classmethod
    def get_file_type(cls, filename: str) -> str:
        """Get file type from filename."""
        ext = filename.lower().split('.')[-1] if '.' in filename else ''
        return ext


# Singleton instance
extractor = FileExtractor()
